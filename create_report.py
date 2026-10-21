"""Tạo báo cáo Word cho hai phần Dữ liệu và Kết quả từ các kết quả đã chạy."""

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from config import DATA_DIR, FIGURES_DIR, REPORT_DIR, RESULTS_DIR


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_table(table, header=True, font_size=8.5):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
            if header and row_index == 0:
                shade_cell(cell, "1F4E78")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.bold = True
                        run.font.name = "Arial"
                        run.font.size = Pt(font_size)
            else:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(font_size)


def add_simple_table(doc, headers, rows, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=(255, 255, 255), size=font_size)
        shade_cell(table.rows[0].cells[i], "1F4E78")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
    format_table(table, header=True, font_size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_figure(doc, filename, caption, width=6.25):
    path = FIGURES_DIR / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(path), width=Inches(width))
    caption_p = doc.add_paragraph(caption)
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(5)
    for run in caption_p.runs:
        run.italic = True
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(89, 89, 89)


def add_side_by_side_figures(doc, items, width=3.05):
    table = doc.add_table(rows=1, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, (filename, caption) in zip(table.rows[0].cells, items):
        cell.text = ""
        path = FIGURES_DIR / filename
        if path.exists():
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=Inches(width))
            cp = cell.add_paragraph(caption)
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(1)
            for run in cp.runs:
                run.italic = True
                run.font.name = "Arial"
                run.font.size = Pt(8)
    # Ẩn đường viền của bảng chỉ dùng để bố trí hình.
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tbl_pr.append(borders)


def main():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    metadata = json.loads((DATA_DIR / "generation_metadata.json").read_text(encoding="utf-8"))
    data_summary = json.loads((RESULTS_DIR / "data_summary.json").read_text(encoding="utf-8"))
    model_summary = json.loads((RESULTS_DIR / "model_summary.json").read_text(encoding="utf-8"))
    metrics = pd.read_csv(RESULTS_DIR / "metrics.csv")
    cycles = pd.read_csv(DATA_DIR / "cycle_dataset.csv")
    users = pd.read_csv(DATA_DIR / "users.csv")
    tuning = pd.read_csv(RESULTS_DIR / "hyperparameter_tuning.csv")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    for style_name, size in [("Title", 22), ("Heading 1", 15), ("Heading 2", 11.5)]:
        styles[style_name].font.name = "Arial"
        styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        styles[style_name].font.size = Pt(size)
        styles[style_name].font.bold = True
        styles[style_name].font.color.rgb = RGBColor(31, 78, 121)

    # Trang 1: tiêu đề và tóm tắt dữ liệu.
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(35)
    title.paragraph_format.space_after = Pt(5)
    run = title.add_run("HỆ THỐNG DỰ BÁO VÀ CẢNH BÁO\nDÒNG TIỀN CÁ NHÂN")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(21)
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = doc.add_paragraph("BÁO CÁO PHẦN DỮ LIỆU VÀ KẾT QUẢ THỬ NGHIỆM")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    for run in subtitle.runs:
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(12)

    add_body(doc, "Mục tiêu: xây dựng dữ liệu dòng tiền giả lập, xác định các chu kỳ có nguy cơ thiếu tiền và so sánh mô hình học máy với các phương pháp cơ sở.")
    add_body(doc, "Lưu ý phạm vi: toàn bộ dữ liệu trong báo cáo là dữ liệu tổng hợp được tạo bằng chương trình, không phải dữ liệu ngân hàng hoặc khảo sát người thật. Các chỉ số hiệu quả chỉ phản ánh bộ dữ liệu và giả định mô phỏng này.")

    doc.add_heading("Tóm tắt quy mô", level=1)
    summary_rows = [
        ("Người dùng giả lập", f"{data_summary['n_users']:,}"),
        ("Thời gian dữ liệu ngày", f"{data_summary['daily_start']} đến {data_summary['daily_end']} ({data_summary['daily_dates']} ngày)"),
        ("Dòng trong daily.csv", f"{data_summary['daily_rows']:,}"),
        ("Mẫu chu kỳ dùng cho mô hình", f"{data_summary['cycle_samples']:,}"),
        ("Nhãn thiếu tiền", f"{data_summary['cycle_shortfall_count']:,}/{data_summary['cycle_samples']:,} ({data_summary['cycle_shortfall_rate']:.2%})"),
        ("Chia tập theo người dùng", f"Train {model_summary['n_train_users']} | Validation {model_summary['n_validation_users']} | Test {model_summary['n_test_users']}"),
        ("Mô hình chính được chọn", model_summary["model_name"]),
    ]
    add_simple_table(doc, ["Hạng mục", "Giá trị"], summary_rows, font_size=9)
    doc.add_page_break()

    # Trang 2: mô tả dữ liệu.
    doc.add_heading("1. Dữ liệu", level=1)
    add_body(doc, f"Bộ dữ liệu ngày mô phỏng {data_summary['n_users']} người trong khoảng {data_summary['daily_start']}–{data_summary['daily_end']}, tổng cộng {data_summary['daily_rows']:,} dòng. Thu nhập hàng tháng được lấy trong khoảng 10–30 triệu đồng. Mỗi hồ sơ có một trong bốn kiểu chi tiêu và một trong hai lịch nhận thu nhập; tám tổ hợp được phân bố gần đồng đều.")
    add_body(doc, "Bốn kiểu chi tiêu gồm: (1) tương đối đều; (2) dồn chi vào đầu chu kỳ; (3) chi tiêu tăng vào cuối tuần; (4) có khoản chi lớn bất thường. Chi tiêu nền có nhiễu log-normal; với nhóm khoản chi bất thường, chương trình mô phỏng thêm các khoản tăng đột biến. Hệ số tổng chi được hiệu chỉnh để tỷ lệ nhãn thiếu tiền của các chu kỳ đầy đủ gần mục tiêu 25%.")
    add_body(doc, f"Hệ số chi tiêu cuối cùng do chương trình hiệu chỉnh là {metadata['spend_scale']:.4f}; tỷ lệ nhãn thiếu tiền trên các chu kỳ hiệu chỉnh là {metadata['actual_shortfall_rate_on_complete_cycles']:.2%}.")

    doc.add_heading("Các tệp dữ liệu", level=2)
    add_simple_table(doc, ["Tệp", "Nội dung"], [
        ("daily.csv", "Một dòng/người/ngày; có tám cột theo yêu cầu ban đầu."),
        ("users.csv", "Thông tin hồ sơ: kiểu chi tiêu, lịch nhận thu nhập và thu nhập tháng."),
        ("cycle_dataset.csv", "Một dòng/chu kỳ thu nhập; chứa đặc trưng đầu chu kỳ và nhãn shortfall."),
    ], font_size=8.5)

    doc.add_heading("Ý nghĩa cột trong daily.csv", level=2)
    column_rows = [
        ("user_id", "Mã người dùng giả lập."),
        ("date", "Ngày ghi nhận."),
        ("spend", "Chi tiêu trong ngày (VND)."),
        ("income", "Thu nhập nhận trong ngày; 0 nếu không có."),
        ("balance", "Số dư cuối ngày sau thu và chi."),
        ("days_to_next_income", "Số ngày đến lần thu nhập tiếp theo."),
        ("day_in_cycle", "Ngày thứ mấy của chu kỳ hiện tại, bắt đầu từ 1."),
        ("pattern", "Kiểu chi tiêu gán cho người dùng."),
    ]
    add_simple_table(doc, ["Cột", "Ý nghĩa"], column_rows, font_size=8)
    doc.add_page_break()

    # Trang 3: nhãn và quy trình tạo mẫu.
    doc.add_heading("1.1. Tạo nhãn và đặc trưng", level=1)
    add_body(doc, f"Từ {data_summary['daily_rows']:,} bản ghi ngày, chương trình tạo {data_summary['cycle_samples']:,} mẫu chu kỳ, tương đương {data_summary['cycle_users']} người × 16 chu kỳ đủ điều kiện. Chu kỳ đầu bị loại nếu chưa có đủ 30 ngày lịch sử; chu kỳ cuối bị loại nếu chưa quan sát được đến lần nhận thu nhập kế tiếp.")
    add_body(doc, "Nhãn shortfall bằng 1 nếu số dư mô phỏng thấp hơn 0 ít nhất một ngày từ ngày nhận thu nhập hiện tại đến ngày trước kỳ thu nhập kế tiếp; bằng 0 nếu không. Đây là định nghĩa theo mô phỏng, không đồng nghĩa với việc quan sát được một khoản nợ thực tế.")
    add_body(doc, "Đặc trưng dự báo được tính từ thông tin có sẵn trước chu kỳ hiện tại: thu nhập và số dư tại đầu chu kỳ, độ dài chu kỳ, trung bình/độ lệch chuẩn/mức chi cao nhất trong 30 ngày trước đó, tỷ trọng chi cuối tuần, tỷ lệ chi của chu kỳ trước, tình trạng thiếu tiền ở chu kỳ trước và kiểu chi tiêu/lịch thu nhập. Chi tiêu tương lai trong chính chu kỳ được dùng để tạo nhãn, không đưa vào đầu vào của mô hình.")
    add_body(doc, "Giả định đơn giản hóa: nếu số dư cuối kỳ trước âm, phần âm được đưa về 0 trước khi cộng thu nhập mới; phần âm biểu diễn nhu cầu chi chưa đủ tiền chứ không phải nợ/overdraft tiếp tục mang sang kỳ sau. Giả định này cần được thay thế nếu có dữ liệu thực tế.")

    doc.add_heading("1.2. Chia dữ liệu và tránh rò rỉ", level=1)
    add_body(doc, f"Dữ liệu được chia theo user_id, không chia ngẫu nhiên từng dòng: {model_summary['n_train_users']} người train ({model_summary['n_train_samples']:,} mẫu), {model_summary['n_validation_users']} người validation ({model_summary['n_validation_samples']:,} mẫu) và {model_summary['n_test_users']} người test ({model_summary['n_test_samples']:,} mẫu). Như vậy một người không xuất hiện đồng thời trong train và test.")
    add_body(doc, "Validation được dùng để chọn cấu hình mô hình và ngưỡng quyết định theo F1. Tập test chỉ dùng một lần để so sánh kết quả cuối, hạn chế việc điều chỉnh mô hình dựa trên test.")
    doc.add_page_break()

    # Trang 4: kết quả tổng hợp.
    doc.add_heading("2. Kết quả", level=1)
    add_body(doc, f"Tỷ lệ nhãn thiếu tiền trong cycle_dataset.csv là {data_summary['cycle_shortfall_rate']:.2%} ({data_summary['cycle_shortfall_count']:,}/{data_summary['cycle_samples']:,} mẫu). Tập test gồm {model_summary['n_test_samples']:,} chu kỳ của {model_summary['n_test_users']} người chưa xuất hiện trong train hoặc validation.")
    add_body(doc, "Ba phương pháp được so sánh: B1 là luật cảnh báo nếu chi tiêu dự kiến cho cả chu kỳ vượt số dư đầu chu kỳ; B2 là Logistic Regression; mô hình chính là cấu hình cây tăng cường được chọn trên validation.")

    metric_rows = []
    for _, row in metrics.iterrows():
        metric_rows.append((
            row["model"],
            f"{row['accuracy']:.3f}",
            f"{row['precision']:.3f}",
            f"{row['recall']:.3f}",
            f"{row['f1']:.3f}",
            f"{row['roc_auc']:.3f}",
            f"{row['pr_auc']:.3f}",
        ))
    add_simple_table(doc, ["Mô hình", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC", "PR-AUC"], metric_rows, font_size=7.4)

    main_row = metrics.loc[metrics["model"].str.startswith("Mô hình chính")].iloc[0]
    b1_row = metrics.loc[metrics["model"].str.startswith("B1")].iloc[0]
    b2_row = metrics.loc[metrics["model"].str.startswith("B2")].iloc[0]
    add_body(doc, f"Trên tập test, mô hình chính đạt F1 = {main_row['f1']:.3f} và Recall = {main_row['recall']:.3f}; baseline B1 đạt F1 = {b1_row['f1']:.3f}, Recall = {b1_row['recall']:.3f}. Như vậy mô hình chính tăng F1 khoảng {(main_row['f1'] - b1_row['f1'])*100:.1f} điểm phần trăm và Recall khoảng {(main_row['recall'] - b1_row['recall'])*100:.1f} điểm phần trăm so với luật B1.")
    add_body(doc, f"Tuy nhiên, Logistic Regression có ROC-AUC = {b2_row['roc_auc']:.3f}, cao hơn một chút so với mô hình chính ({main_row['roc_auc']:.3f}). Do đó, theo bộ thử nghiệm hiện tại, mô hình chính có F1/Recall tốt hơn, nhưng không đứng đầu ở mọi chỉ số.")
    add_body(doc, f"Cấu hình cây tăng cường được chọn: {model_summary['model_name']}. Ngưỡng cảnh báo chính được chọn trên validation là {model_summary['selected_threshold']:.2f}; không chọn ngưỡng dựa trên tập test.")
    doc.add_page_break()

    # Trang 5: phân bố và so sánh.
    doc.add_heading("2.1. Phân bố nhãn và kiểu chi tiêu", level=1)
    add_body(doc, "Biểu đồ dưới đây cho biết tỷ lệ thiếu tiền toàn bộ dữ liệu và sự khác nhau giữa các kiểu chi tiêu. Các khác biệt phản ánh trực tiếp các giả định được lập trình trong bộ sinh dữ liệu; không được diễn giải như kết luận về hành vi tài chính của người thật.")
    add_side_by_side_figures(doc, [
        ("01_label_distribution.png", "Hình 1. Phân bố hai lớp nhãn."),
        ("02_shortfall_by_pattern.png", "Hình 2. Tỷ lệ thiếu tiền theo kiểu chi tiêu."),
    ], width=3.2)
    doc.add_paragraph()
    add_figure(doc, "03_shortfall_by_income_type.png", "Hình 3. Tỷ lệ thiếu tiền theo lịch nhận thu nhập.", width=5.8)
    doc.add_page_break()

    # Trang 6: metric và ROC.
    doc.add_heading("2.2. So sánh mô hình", level=1)
    add_body(doc, "F1 kết hợp Precision và Recall; Recall quan trọng nếu muốn hạn chế bỏ sót các chu kỳ thiếu tiền, còn Precision thể hiện mức độ đáng tin của cảnh báo. ROC-AUC và PR-AUC đo khả năng xếp hạng rủi ro ở nhiều ngưỡng, không chỉ tại ngưỡng cảnh báo được chọn.")
    add_side_by_side_figures(doc, [
        ("04_model_comparison.png", "Hình 4. So sánh F1, Recall và ROC-AUC."),
        ("05_roc_curve.png", "Hình 5. Đường ROC trên tập test."),
    ], width=3.25)
    doc.add_page_break()

    # Trang 7: giải thích mô hình.
    doc.add_heading("2.3. Đặc trưng và lỗi dự báo", level=1)
    add_body(doc, "Feature importance hỗ trợ xem mô hình dựa vào nhóm thông tin nào nhiều hơn. Đây là mức độ đóng góp trong mô hình đã huấn luyện, không chứng minh quan hệ nhân quả. Ma trận nhầm lẫn cho biết số chu kỳ thiếu tiền được bắt đúng, bị bỏ sót hoặc cảnh báo nhầm.")
    add_side_by_side_figures(doc, [
        ("06_feature_importance.png", "Hình 6. Đặc trưng quan trọng."),
        ("07_confusion_matrix.png", "Hình 7. Ma trận nhầm lẫn của mô hình chính."),
    ], width=3.25)
    doc.add_page_break()

    # Trang 8: đường số dư và tổng thu chi.
    doc.add_heading("2.4. Diễn biến dòng tiền mô phỏng", level=1)
    add_body(doc, "Các ví dụ và tổng hợp theo tháng dưới đây dùng để kiểm tra trực quan dữ liệu mô phỏng. Số dư âm biểu diễn nhu cầu chi vượt nguồn tiền sẵn có theo giả định của mô hình; trong một ứng dụng thực tế, cần phân biệt số dư tài khoản thật với phần chi tiêu chưa được tài trợ.")
    add_figure(doc, "08_balance_examples.png", "Hình 8. Ví dụ diễn biến số dư ở bốn kiểu chi tiêu.", width=6.1)
    add_figure(doc, "09_monthly_income_spend.png", "Hình 9. Tổng thu nhập và chi tiêu theo tháng.", width=6.1)
    doc.add_page_break()

    # Kết luận.
    doc.add_heading("Kết luận và giới hạn", level=1)
    add_body(doc, f"Pipeline đã tạo {data_summary['daily_rows']:,} bản ghi ngày, {data_summary['cycle_samples']:,} mẫu chu kỳ và hoàn tất so sánh ba phương pháp. Trong lần chạy này, mô hình chính đạt F1 = {main_row['f1']:.3f}; kết quả tốt hơn baseline luật theo F1/Recall, nhưng Logistic Regression vẫn có ROC-AUC cao hơn một chút.")
    add_body(doc, "Giới hạn chính: dữ liệu không phải dữ liệu người thật; kiểu chi tiêu, thời điểm thu nhập và mức độ bất thường được quy định từ trước; nhãn được tạo trực tiếp từ số dư mô phỏng. Vì vậy kết quả chỉ chứng minh quy trình kỹ thuật có thể chạy đầu-cuối, chưa chứng minh hiệu quả trên dữ liệu đời thực. Bước tiếp theo nên thử độ nhạy theo nhiều seed và tham số, hiệu chỉnh theo dữ liệu thật được cấp phép, kiểm tra theo thời gian và đánh giá chi phí của cảnh báo nhầm/bỏ sót.")

    # Footer.
    for sec in doc.sections:
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = footer.add_run("Hệ thống Dự báo và Cảnh báo Dòng tiền Cá nhân | Dữ liệu giả lập")
        rr.font.name = "Arial"
        rr.font.size = Pt(8)
        rr.font.color.rgb = RGBColor(100, 100, 100)

    out = REPORT_DIR / "Bao_cao_Du_lieu_Ket_qua.docx"
    doc.save(out)
    print(f"Đã tạo báo cáo: {out}")


if __name__ == "__main__":
    main()
